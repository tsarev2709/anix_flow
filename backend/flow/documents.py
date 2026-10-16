from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

MAX_UPLOAD = 12 * 1024 * 1024
MAX_TEXT = 120_000


def extract(name: str, data: bytes) -> tuple[str, str]:
    extension = Path(name).suffix.lower()
    warning = ""
    if extension in (".txt", ".md", ".csv", ".json"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("Текст должен быть в UTF-8. Пересохрани файл и загрузи его снова.")
    elif extension == ".pdf":
        from pypdf import PdfReader

        try:
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise ValueError("PDF защищён паролем. Нужна доступная текстовая версия.")
            if len(reader.pages) > 300:
                raise ValueError("В PDF больше 300 страниц. Выбери нужный раздел.")
            parts = []
            length = 0
            for index, page in enumerate(reader.pages):
                chunk = f"[Страница {index+1}]\n{page.extract_text() or ''}"
                length += len(chunk)
                if length > MAX_TEXT:
                    raise ValueError("Извлечённый текст длиннее 120 000 символов. Выбери нужный раздел.")
                parts.append(chunk)
            text = "\n\n".join(parts)
            if sum(len((p.extract_text() or "").strip()) for p in reader.pages) < 30:
                raise ValueError("В PDF нет читаемого текста. OCR пока не подключён; загрузи расшифровку.")
            warning = "Проверь порядок текста и таблицы. OCR не выполнялся."
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Не удалось прочитать PDF. Проверь файл или экспортируй текст.") from exc
    elif extension == ".docx":
        from docx import Document

        try:
            with ZipFile(BytesIO(data)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                    raise ValueError("Распакованный DOCX слишком большой.")
            document = Document(BytesIO(data))
            text = "\n".join([p.text for p in document.paragraphs] +
                             [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows])
            warning = "Проверь порядок таблиц. Колонтитулы, комментарии и изображения не извлекаются."
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Не удалось прочитать DOCX. Проверь файл или экспортируй текст.") from exc
    else:
        raise ValueError("Поддерживаются TXT, MD, CSV, JSON, PDF с текстовым слоем и DOCX.")
    if len(text) > MAX_TEXT:
        raise ValueError("Текст длиннее 120 000 символов. Раздели документ; файл не был обрезан.")
    if not text.strip():
        raise ValueError("Файл не содержит текста.")
    return text.strip(), warning
