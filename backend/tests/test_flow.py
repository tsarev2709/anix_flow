from io import BytesIO

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pydantic import ValidationError
from pypdf import PdfWriter

from flow import llm
from flow.documents import extract
from flow.engine import evidence_issues, fingerprint, review_variant, template_spec, template_variants, words
from flow.main import app
from flow.models import ProjectInput, Source, Spec


@pytest.fixture
def project():
    return ProjectInput(id="test-project", title="Первый день", task="Помочь обратиться к наставнику",
                        audience="Новые сотрудники", before="Стесняются спросить", desired="Уверены в поддержке",
                        action="Напиши наставнику свой вопрос.", success="Может назвать действие",
                        claims=[{"id": "C1", "text": "Задавать вопросы совершенно нормально."}],
                        accents=[{"id": "A1", "object": "Обращение", "claim_id": "C1", "intent": "normalize"}])


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("FLOW_DATA_DIR", str(tmp_path))
    with TestClient(app) as session:
        yield session


def variant(project):
    return template_variants(project, template_spec(project)).variants[0]


def test_approval_and_staleness(client, project):
    body = {"project": project.model_dump(), "mode": "template"}
    brief = client.post("/api/flow/brief", json=body)
    assert brief.status_code == 200
    info = brief.json()
    request = {**body, "spec": info["spec"], "fingerprint": info["fingerprint"]}
    assert client.post("/api/flow/variants", json=request).status_code == 409
    request["approved"] = True
    result = client.post("/api/flow/variants", json=request)
    assert result.status_code == 200
    assert len(result.json()["variants"]) == 3
    assert result.json()["mode"] == "template"
    request["project"]["tone"]["warmth"] = 90
    assert client.post("/api/flow/variants", json=request).status_code == 409


def test_fingerprint_tracks_sources_and_order(project):
    copy = ProjectInput.model_validate(project.model_dump())
    assert fingerprint(project) == fingerprint(copy)
    copy.desired = "Другое изменение"
    assert fingerprint(project) != fingerprint(copy)


def test_nonoverlapping_timing_and_words(project):
    v = variant(project)
    report = review_variant(project, v)
    assert report["duration"] == 45
    assert report["failures"] == 0
    v.scenes[1].speech = " ".join(["слово"] * 50) + "."
    report = review_variant(project, v)
    assert report["bad_sentences"][0]["words"] == 50
    assert any(ch["code"] == "scene_speech" and ch["status"] == "fail" for ch in report["checks"])


def test_minimum_sentence_words_and_russian_count(project):
    assert words("Кто-то произнёс: «Это всё-таки 2026!»") == ["Кто-то", "произнёс", "Это", "всё-таки", "2026"]
    v = variant(project)
    v.scenes[0].speech = "Да!"
    report = review_variant(project, v)
    assert report["bad_sentences"][0]["words"] == 1


def test_pauses_and_duration_are_distinct_checks(project):
    v = variant(project)
    v.scenes[0].duration = 1
    report = review_variant(project, v)
    assert any(ch["code"] == "duration" and ch["status"] == "fail" for ch in report["checks"])
    v.scenes[0].duration = 9
    project.constraints.silent_seconds = 44
    report = review_variant(project, v)
    assert any(ch["code"] == "speech_budget" and ch["status"] == "fail" for ch in report["checks"])


def test_required_phrases_must_be_spoken_and_forbidden_in_visual(project):
    v = variant(project)
    project.constraints.must_include = "Особая обязательная фраза"
    project.constraints.forbidden = "запрещённая фраза"
    v.scenes[0].visual = "Особая обязательная фраза и запрещённая фраза"
    checks = review_variant(project, v)["checks"]
    assert any(ch["code"] == "must_include" and ch["status"] == "fail" for ch in checks)
    assert any(ch["code"] == "forbidden" and ch["status"] == "fail" for ch in checks)


def test_unknown_scene_links_and_duplicate_ids(project):
    v = variant(project)
    v.scenes[0].claim_ids = ["made-up"]
    v.scenes[0].accent_ids = ["fake-accent"]
    v.scenes[0].id = v.scenes[1].id
    coverage = next(ch for ch in review_variant(project, v)["checks"] if ch["code"] == "coverage")
    assert coverage["status"] == "fail"
    assert "made-up" in coverage["detail"] and "fake-accent" in coverage["detail"]


def test_claim_evidence_must_be_literal(project):
    project.sources = [Source(id="s1", name="Бриф", text="Задавать вопросы нормально.")]
    project.claims[0].source_id = "s1"
    project.claims[0].quote = "Совсем другая цитата"
    assert evidence_issues(project)
    project.claims[0].quote = "Задавать вопросы нормально."
    assert not evidence_issues(project)
    project.claims[0].origin = "inferred"
    assert evidence_issues(project)
    project.claims[0].confirmed = True
    assert not evidence_issues(project)


def test_unconfirmed_claim_blocks_generation(client, project):
    project.claims[0].origin = "inferred"
    body = {"project": project.model_dump(), "mode": "template", "spec": template_spec(project).model_dump(),
            "fingerprint": fingerprint(project), "approved": True}
    assert client.post("/api/flow/variants", json=body).status_code == 409


@pytest.mark.parametrize("changes", [{"min_words": 20, "max_words": 10}, {"silent_seconds": 45}, {"wpm": 0}, {"duration": -1}, {"variants": 5}])
def test_invalid_constraints_rejected(client, project, changes):
    body = project.model_dump()
    body["constraints"].update(changes)
    assert client.post("/api/flow/brief", json={"project": body}).status_code == 422


def test_duplicate_semantic_ids_rejected(project):
    body = project.model_dump()
    body["claims"].append(body["claims"][0])
    with pytest.raises(ValidationError):
        ProjectInput.model_validate(body)


def test_upload_utf8_docx_and_scan(client):
    response = client.post("/api/flow/documents", files={"file": ("brief.txt", "Задача клиента".encode())})
    assert response.status_code == 200
    assert response.json()["text"] == "Задача клиента"
    assert client.post("/api/flow/documents", files={"file": ("bad.txt", b"\xff\xfe")}).status_code == 422
    document = Document(); document.add_paragraph("Важный тезис")
    table = document.add_table(rows=1, cols=2); table.cell(0, 0).text = "Цель"; table.cell(0, 1).text = "Действие"
    buffer = BytesIO(); document.save(buffer)
    response = client.post("/api/flow/documents", files={"file": ("brief.docx", buffer.getvalue())})
    assert response.status_code == 200 and "Цель | Действие" in response.json()["text"]
    pdf = PdfWriter(); pdf.add_blank_page(width=100, height=100); buffer = BytesIO(); pdf.write(buffer)
    assert client.post("/api/flow/documents", files={"file": ("scan.pdf", buffer.getvalue())}).status_code == 422
    assert client.post("/api/flow/documents", files={"file": ("image.png", b"data")}).status_code == 422


def test_context_does_not_silently_truncate(client, project, monkeypatch):
    monkeypatch.setenv("FLOW_MAX_CONTEXT_CHARS", "10")
    response = client.post("/api/flow/brief", json={"project": project.model_dump(), "mode": "template"})
    assert response.status_code == 413


def test_sqlite_backups_roundtrip(client, project):
    data = {"id": project.id, "data": project.model_dump()}
    assert client.put(f"/api/flow/projects/{project.id}", json=data).status_code == 200
    records = client.get("/api/flow/projects").json()["projects"]
    assert records[0]["data"]["task"] == project.task
    assert client.put("/api/flow/projects/different", json=data).status_code == 422


def test_untrusted_origin_cannot_mutate(client, project):
    response = client.post("/api/flow/brief", headers={"Origin": "https://untrusted.example"},
                           json={"project": project.model_dump(), "mode": "template"})
    assert response.status_code == 403
    response = client.options("/api/flow/brief", headers={"Origin": "http://127.0.0.1:3000", "Access-Control-Request-Method": "POST"})
    assert response.status_code == 200


def test_model_is_not_silently_replaced_by_templates(client, project, monkeypatch):
    async def fail(*args, **kwargs):
        raise llm.ModelError("Модель недоступна")
    monkeypatch.setattr(llm, "generate", fail)
    response = client.post("/api/flow/brief", json={"project": project.model_dump()})
    assert response.status_code == 502 and "Модель недоступна" in response.json()["detail"]


def test_hallucinated_quote_rejected(client, project, monkeypatch):
    async def bad_quote(*args, **kwargs):
        spec = template_spec(project)
        spec.evidence[0].source_id = "missing"
        spec.evidence[0].quote = "invented"
        return spec
    monkeypatch.setattr(llm, "generate", bad_quote)
    assert client.post("/api/flow/brief", json={"project": project.model_dump()}).status_code == 502


@pytest.mark.asyncio
async def test_ollama_schema_and_repair(project, monkeypatch):
    requests = []
    valid = template_spec(project).model_dump_json()
    def handler(request):
        import json
        requests.append(json.loads(request.content))
        content = '{"broken": true}' if len(requests) == 1 else valid
        return httpx.Response(200, json={"message": {"content": content}})
    real_client = httpx.AsyncClient
    monkeypatch.setattr(llm.httpx, "AsyncClient", lambda **kw: real_client(transport=httpx.MockTransport(handler), **kw))
    monkeypatch.setenv("FLOW_LLM_PROVIDER", "ollama")
    result = await llm.generate(Spec, "brief")
    assert result.main_message == project.claims[0].text
    assert len(requests) == 2
    assert requests[0]["format"]["type"] == "object" and requests[0]["think"] is False


@pytest.mark.asyncio
async def test_compatible_provider_uses_structured_output(project, monkeypatch):
    def handler(request):
        import json
        body = json.loads(request.content)
        assert request.url.path == "/v1/chat/completions"
        assert body["response_format"]["type"] == "json_schema"
        return httpx.Response(200, json={"choices": [{"message": {"content": template_spec(project).model_dump_json()}}]})
    real_client = httpx.AsyncClient
    monkeypatch.setattr(llm.httpx, "AsyncClient", lambda **kw: real_client(transport=httpx.MockTransport(handler), **kw))
    monkeypatch.setenv("FLOW_LLM_PROVIDER", "compatible")
    monkeypatch.setenv("FLOW_LLM_BASE_URL", "http://localhost:1234/v1")
    assert (await llm.generate(Spec, "brief")).viewer_action == project.action


def test_diagnostics_prioritize_conflicts_and_preserve_unknown(client, project):
    project.before = ""
    project.constraints.must_include = "Обязательная запрещённая фраза"
    project.constraints.forbidden = "запрещённая фраза"
    response = client.post("/api/flow/diagnostics", json=project.model_dump())
    assert response.status_code == 200
    data = response.json()
    assert data["findings"][0]["code"] == "phrase_conflict"
    assert data["blockers"] == 1
    assert any(f["code"] == "missing_before" for f in data["findings"])
    assert data["fingerprint"] == fingerprint(project)
    assert project.before == ""


def test_diagnostics_does_not_require_behavior_for_emotional_work(client, project):
    project.effects = ["emotion"]
    project.action = ""
    findings = client.post("/api/flow/diagnostics", json=project.model_dump()).json()["findings"]
    assert not any(f["code"] == "missing_action" for f in findings)
    project.effects = ["action"]
    findings = client.post("/api/flow/diagnostics", json=project.model_dump()).json()["findings"]
    assert any(f["code"] == "missing_action" for f in findings)


def test_impossible_mandatory_speech_blocks_before_model_call(client, project, monkeypatch):
    async def must_not_call(*args, **kwargs):
        pytest.fail("Contradictions must be found before a model request")
    monkeypatch.setattr(llm, "generate", must_not_call)
    project.constraints.must_include = " ".join(["слово"] * 200)
    response = client.post("/api/flow/variants", json={"project": project.model_dump(),
        "spec": template_spec(project).model_dump(), "approved": True, "fingerprint": fingerprint(project)})
    assert response.status_code == 409
    assert "200" in response.json()["detail"]


def test_required_fragments_are_not_naively_summed(client, project):
    project.constraints.duration = 10
    project.constraints.tolerance = 0
    project.constraints.silent_seconds = 1
    phrase = " ".join(["слово"] * 15)
    project.constraints.must_include = phrase + "\n" + phrase.upper() + "\n" + " ".join(["слово"] * 14)
    data = client.post("/api/flow/diagnostics", json=project.model_dump()).json()
    assert data["required_words_lower_bound"] == 15
    assert not data["blockers"]


def test_missing_accent_is_reported(project):
    v = variant(project)
    for scene in v.scenes:
        scene.accent_ids = []
    check = next(c for c in review_variant(project, v)["checks"] if c["code"] == "coverage")
    assert check["status"] == "fail" and "A1" in check["detail"]


def test_context_limit_counts_approved_spec(client, project, monkeypatch):
    from flow.engine import prompt_for
    spec = template_spec(project)
    limit = len(prompt_for(project, "brief")) + 50
    monkeypatch.setenv("FLOW_MAX_CONTEXT_CHARS", str(limit))
    spec.rationale = "длинное обоснование " * 500
    response = client.post("/api/flow/variants", json={"project": project.model_dump(), "mode": "template",
        "spec": spec.model_dump(), "fingerprint": fingerprint(project), "approved": True})
    assert response.status_code == 413


def test_context_limit_counts_repair_payload(client, project, monkeypatch):
    from flow.engine import prompt_for
    spec = template_spec(project)
    monkeypatch.setenv("FLOW_MAX_CONTEXT_CHARS", str(len(prompt_for(project, "variants", spec)) + 100))
    calls = []
    async def invalid_result(*args, **kwargs):
        calls.append(args)
        result = template_variants(project, spec)
        result.variants[0].scenes[0].duration = 1
        return result
    monkeypatch.setattr(llm, "generate", invalid_result)
    response = client.post("/api/flow/variants", json={"project": project.model_dump(),
        "spec": spec.model_dump(), "fingerprint": fingerprint(project), "approved": True})
    assert response.status_code == 413
    assert len(calls) == 1


def test_backups_reject_stale_writes_and_preserve_revisions(client, project):
    path = f"/api/flow/projects/{project.id}"
    body = {"id": project.id, "data": {"title": "Первая версия"}}
    saved = client.put(path, json=body).json()
    assert saved["revision"] == 1
    assert client.put(path, json=body).status_code == 409
    body.update(expected_revision=1, data={"title": "Вторая версия"})
    assert client.put(path, json=body).json()["revision"] == 2
    body["data"]["title"] = "Устаревшая запись"
    assert client.put(path, json=body).status_code == 409
    assert client.get("/api/flow/projects").json()["projects"][0]["data"]["title"] == "Вторая версия"
    revisions = client.get(path + "/revisions").json()["revisions"]
    assert [r["revision"] for r in revisions] == [2, 1]
    assert revisions[-1]["data"]["title"] == "Первая версия"


def test_old_sqlite_database_upgrades_without_data_loss(client, project, tmp_path):
    import json
    import sqlite3
    with sqlite3.connect(tmp_path / "projects.sqlite3") as db:
        db.execute("CREATE TABLE projects (id TEXT PRIMARY KEY, updated TEXT NOT NULL, data TEXT NOT NULL)")
        db.execute("INSERT INTO projects VALUES (?, ?, ?)", (project.id, "2026-10-15", json.dumps({"title": "Старая копия"})))
    records = client.get("/api/flow/projects").json()["projects"]
    assert records[0]["revision"] == 1
    assert records[0]["data"]["title"] == "Старая копия"
    saved = client.put(f"/api/flow/projects/{project.id}", json={"id": project.id, "data": {"title": "Новая"}, "expected_revision": 1})
    assert saved.json()["revision"] == 2
    assert len(client.get(f"/api/flow/projects/{project.id}/revisions").json()["revisions"]) == 2


def test_concurrent_backups_have_one_winner(tmp_path, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from flow import storage
    monkeypatch.setenv("FLOW_DATA_DIR", str(tmp_path))
    storage.save("race", {"value": "initial"})
    def write(value):
        try:
            storage.save("race", {"value": value}, expected_revision=1)
            return "saved"
        except storage.ConflictError:
            return "conflict"
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(write, ["a", "b"]))
    assert sorted(results) == ["conflict", "saved"]
    assert len(storage.revisions("race")) == 2
